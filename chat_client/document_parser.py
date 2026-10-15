"""
@input: os, logging, pathlib; pypdf, docx, pptx, openpyxl, bs4; rapidocr(可选)
@output: parse_document, chunk_text, is_supported, extract_text_from_image, SUPPORTED_EXTS
@position: Knowledge layer — 文档解析 + 切块 + OCR，供知识库入库管线使用
@auto-doc: Update header and folder INDEX.md when this file changes

参考 TencentCloud/Octop 的 knowledge/parse.py + chunk.py 设计，
按 BT 现有依赖（docx/pptx/openpyxl/bs4 + rapidocr）实现。
"""
from __future__ import annotations

import html as _html
import io
import logging
import os
import re
import zipfile
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# 支持的文本类文档后缀
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".log", ".json", ".yaml", ".yml", ".ini", ".conf"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
HTML_EXTS = {".html", ".htm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}
LEGACY_EXTS = {".doc", ".ppt", ".xls"}  # 旧二进制格式，不支持，给出提示

SUPPORTED_EXTS = (TEXT_EXTS | PDF_EXTS | DOCX_EXTS | PPTX_EXTS
                  | XLSX_EXTS | HTML_EXTS | IMAGE_EXTS)


def is_supported(path: str | os.PathLike) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_EXTS


def _read_text_file(path: Path) -> str:
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
        except Exception:
            break
    # 最后兜底：二进制安全读取
    try:
        return path.read_bytes().decode("utf-8", errors="replace")
    except Exception:
        return ""


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[解析失败] 缺少 pypdf，请执行: .venv/bin/pip install pypdf"
    try:
        reader = PdfReader(str(path))
        parts = []
        for i, page in enumerate(reader.pages):
            try:
                txt = page.extract_text() or ""
            except Exception as e:
                logger.debug("PDF 第%d页解析失败: %s", i, e)
                txt = ""
            if txt.strip():
                parts.append(txt)
        return "\n\n".join(parts)
    except Exception as e:
        logger.warning("PDF 解析异常 %s: %s", path, e)
        return f"[PDF 解析失败] {e}"


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return "[解析失败] 缺少 python-docx"
    try:
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # 表格
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        logger.warning("DOCX 解析异常 %s: %s", path, e)
        return f"[DOCX 解析失败] {e}"


def _parse_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "[解析失败] 缺少 python-pptx"
    try:
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_parts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    slide_parts.append(shape.text_frame.text)
            if slide_parts:
                parts.append(f"# 第{i}页\n" + "\n".join(slide_parts))
        return "\n\n".join(parts)
    except Exception as e:
        logger.warning("PPTX 解析异常 %s: %s", path, e)
        return f"[PPTX 解析失败] {e}"


def _parse_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return "[解析失败] 缺少 openpyxl"
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows_txt = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    rows_txt.append(" | ".join(cells))
            if rows_txt:
                parts.append(f"## 工作表: {ws.title}\n" + "\n".join(rows_txt))
        wb.close()
        return "\n\n".join(parts)
    except Exception as e:
        logger.warning("XLSX 解析异常 %s: %s", path, e)
        return f"[XLSX 解析失败] {e}"


def _parse_html(path: Path) -> str:
    raw = _read_text_file(path)
    return _strip_html(raw)


def _strip_html(raw: str) -> str:
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return soup.get_text(separator="\n", strip=True)
    except Exception:
        # 兜底正则去标签
        txt = re.sub(r"<script.*?</script>", "", raw, flags=re.S | re.I)
        txt = re.sub(r"<style.*?</style>", "", txt, flags=re.S | re.I)
        txt = re.sub(r"<[^>]+>", " ", txt)
        return _html.unescape(re.sub(r"\s+", " ", txt)).strip()


# ---------------- OCR ----------------

_ocr_engine = None
_ocr_failed = False


def _get_ocr_engine():
    """懒加载 rapidocr 单例（本地离线，中英文）。不可用时返回 None。"""
    global _ocr_engine, _ocr_failed
    if _ocr_engine is not None or _ocr_failed:
        return _ocr_engine
    try:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
        logger.info("[Knowledge] rapidocr 引擎已加载")
    except Exception as e:
        _ocr_failed = True
        logger.warning("[Knowledge] rapidocr 不可用，OCR 降级关闭: %s", e)
    return _ocr_engine


def ocr_available() -> bool:
    return _get_ocr_engine() is not None


def extract_text_from_image(path: str | os.PathLike, threshold: float = 0.5) -> str:
    """对单张图片做 OCR，返回按阅读顺序拼接的文本。"""
    engine = _get_ocr_engine()
    if engine is None:
        return ""
    try:
        result, _ = engine(str(path), use_det=True, use_cls=True, use_rec=True)
    except Exception as e:
        logger.warning("OCR 识别失败 %s: %s", path, e)
        return ""
    if not result:
        return ""
    lines = []
    for item in result:
        try:
            box, text, conf = item[0], item[1], item[2]
        except Exception:
            continue
        if conf is not None and conf < threshold:
            continue
        y = int(box[0][1]) if box else 0
        x = int(box[0][0]) if box else 0
        lines.append((y // 10 * 10, x, text))
    lines.sort(key=lambda t: (t[0], t[1]))
    return "\n".join(t[2] for t in lines)


def _extract_pdf_images_text(path: Path) -> str:
    """提取 PDF 内嵌图片并 OCR（扫描件场景）。"""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
    except Exception:
        return ""
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            for img in page.images:
                data = img.data if hasattr(img, "data") else None
                if not data:
                    continue
                text = _ocr_bytes(data, name=getattr(img, "name", f"p{i}"))
                if text.strip():
                    parts.append(f"[第{i + 1}页图片OCR]\n{text}")
        except Exception as e:
            logger.debug("PDF 图片提取失败 p%d: %s", i, e)
    return "\n\n".join(parts)


def _extract_docx_images_text(path: Path) -> str:
    """提取 DOCX 内嵌图片并 OCR。"""
    parts = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.startswith("word/media/") and Path(name).suffix.lower() in IMAGE_EXTS:
                    text = _ocr_bytes(zf.read(name), name=name)
                    if text.strip():
                        parts.append(f"[文档内嵌图片OCR: {Path(name).name}]\n{text}")
    except Exception as e:
        logger.debug("DOCX 图片提取失败: %s", e)
    return "\n\n".join(parts)


def _extract_pptx_images_text(path: Path) -> str:
    parts = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.startswith("ppt/media/") and Path(name).suffix.lower() in IMAGE_EXTS:
                    text = _ocr_bytes(zf.read(name), name=name)
                    if text.strip():
                        parts.append(f"[幻灯片内嵌图片OCR: {Path(name).name}]\n{text}")
    except Exception as e:
        logger.debug("PPTX 图片提取失败: %s", e)
    return "\n\n".join(parts)


def _ocr_bytes(data: bytes, name: str = "image") -> str:
    """对图片字节流 OCR：写临时文件后识别（rapidocr 接受路径/ndarray）。"""
    engine = _get_ocr_engine()
    if engine is None or not data:
        return ""
    tmp = None
    try:
        import tempfile
        suffix = Path(name).suffix or ".png"
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as f:
            f.write(data)
            tmp = f.name
        result, _ = engine(tmp, use_det=True, use_cls=True, use_rec=True)
        if not result:
            return ""
        lines = []
        for item in result:
            try:
                box, text, conf = item[0], item[1], item[2]
            except Exception:
                continue
            if conf is not None and conf < 0.5:
                continue
            y = int(box[0][1]) if box else 0
            x = int(box[0][0]) if box else 0
            lines.append((y // 10 * 10, x, text))
        lines.sort(key=lambda t: (t[0], t[1]))
        return "\n".join(t[2] for t in lines)
    except Exception as e:
        logger.debug("OCR 字节流失败 %s: %s", name, e)
        return ""
    finally:
        if tmp:
            try:
                os.remove(tmp)
            except Exception:
                pass


# ---------------- 主入口 ----------------

def parse_document(path: str | os.PathLike, *, ocr: bool = True) -> dict[str, Any]:
    """解析文档为纯文本 + 元信息。

    Returns:
        {"text": str, "ext": str, "ocr_used": bool, "warnings": [str], "chars": int}
    """
    p = Path(path)
    ext = p.suffix.lower()
    warnings: list[str] = []
    ocr_used = False
    text = ""

    if not p.exists():
        return {"text": "", "ext": ext, "ocr_used": False,
                "warnings": [f"文件不存在: {p}"], "chars": 0}

    if ext in TEXT_EXTS:
        text = _read_text_file(p)
    elif ext == ".pdf":
        text = _parse_pdf(p)
        if ocr and len(text.strip()) < 50:
            img_txt = _extract_pdf_images_text(p)
            if img_txt.strip():
                text = (text + "\n\n" + img_txt).strip()
                ocr_used = True
            elif not text.strip():
                warnings.append("PDF 无文本层且 OCR 无结果（可能扫描件或加密）")
    elif ext == ".docx":
        text = _parse_docx(p)
        if ocr:
            img_txt = _extract_docx_images_text(p)
            if img_txt.strip():
                text = (text + "\n\n" + img_txt).strip()
                ocr_used = True
    elif ext == ".pptx":
        text = _parse_pptx(p)
        if ocr:
            img_txt = _extract_pptx_images_text(p)
            if img_txt.strip():
                text = (text + "\n\n" + img_txt).strip()
                ocr_used = True
    elif ext in XLSX_EXTS:
        text = _parse_xlsx(p)
    elif ext in HTML_EXTS:
        text = _parse_html(p)
    elif ext in IMAGE_EXTS:
        text = extract_text_from_image(p)
        ocr_used = bool(text.strip())
        if not text.strip():
            warnings.append("图片 OCR 无识别结果（或 rapidocr 不可用）")
    elif ext in LEGACY_EXTS:
        warnings.append(f"不支持旧格式 {ext}，请先转换为 .docx/.pptx/.xlsx")
    else:
        # 未知扩展名：尝试按文本读
        text = _read_text_file(p)
        if not text.strip():
            warnings.append(f"不支持的格式: {ext}")

    text = _normalize(text)
    return {"text": text, "ext": ext, "ocr_used": ocr_used,
            "warnings": warnings, "chars": len(text)}


def _normalize(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def chunk_text(text: str, *, size: int = 800, overlap: int = 120) -> list[str]:
    """固定窗口重叠切块（字符数），与 Octop knowledge/chunk.py 对齐。

    Args:
        size: 单块最大字符数
        overlap: 相邻块重叠字符数（须 < size）
    """
    if not text:
        return []
    if overlap < 0 or overlap >= size:
        raise ValueError("overlap 必须 >=0 且 < size")
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    step = size - overlap
    chunks = []
    i = 0
    n = len(text)
    while i < n:
        chunk = text[i:i + size].strip()
        if chunk:
            chunks.append(chunk)
        if i + size >= n:
            break
        i += step
    return chunks
